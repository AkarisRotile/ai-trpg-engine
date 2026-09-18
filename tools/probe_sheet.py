"""探查 COC7 空白卡 xlsx 的结构，找出关键字段落在哪些单元格。

用法：<venv>\\Scripts\\python.exe tools\\probe_sheet.py
结果写到 data\\sheet_map.txt（UTF-8），避免控制台编码捣乱。
"""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import openpyxl  # noqa: E402

SRC = ROOT / "COC7空白卡CY22.4 Plus.xlsx"
OUT = ROOT / "data" / "sheet_map.txt"

KEYS = [
    "力量", "STR", "敏捷", "DEX", "意志", "POW", "体质", "CON",
    "外貌", "APP", "教育", "EDU", "体型", "SIZ", "智力", "INT",
    "理智", "SAN", "幸运", "LUCK", "生命", "HP", "魔法", "MP",
    "伤害加值", "DB", "移动", "MOV", "母语", "闪避", "职业", "年龄",
    "姓名", "玩家", "信用", "技能", "背景", "调查员",
]


def main() -> int:
    wb = openpyxl.load_workbook(SRC, data_only=False)
    lines: list[str] = [f"文件：{SRC.name}", f"工作表：{wb.sheetnames}", ""]

    for name in wb.sheetnames:
        ws = wb[name]
        hits: list[str] = []
        for row in ws.iter_rows():
            for c in row:
                v = c.value
                if v is None or not isinstance(v, str):
                    continue
                s = v.strip()
                if not s or len(s) > 24:
                    continue
                for k in KEYS:
                    if k in s:
                        hits.append(f"  {c.coordinate:>6}  {s!r}")
                        break
                if len(hits) > 120:
                    break
            if len(hits) > 120:
                break
        lines.append(f"===== {name}  ({ws.max_row} 行 × {ws.max_column} 列) =====")
        lines.extend(hits[:120] or ["  （没有命中关键字段）"])
        lines.append("")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print(f"已写入 {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
