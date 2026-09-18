"""多格式文档读取：把各种乱七八糟的模组文件读成纯文本。

现实里从网上找到的模组长什么样？看看这个真实例子：

    不更文-吞噬深渊之影1+2/
        永远的薄暮之国+疫病之村本体.doc      ← 1 MB 的老式 Word 二进制
        第一章文件/地图带字.png
        第二章文件/怪物资料.docx
        第二章文件/薄暮之国时间线.xls         ← 老式 Excel 二进制
        第二章文件/附魔&商店&咒文&支线任务&粉尘摄入表.xlsx

也就是说：`.doc` / `.docx` / `.xls` / `.xlsx` / `.png` 全都有，一个 `.md` 都没有。
只认 Markdown 的解析器在这种文件夹面前等于瞎了。

所以这里按格式分别处理：

    .docx  python-docx，能认出标题层级 → 输出带 `## ` 的文本，切场景就有依据
    .xlsx  openpyxl，一个工作表一段，行内用 ` | ` 分隔
    .doc   olefile 取 WordDocument 流，按 UTF-16LE 扫可读文本段（启发式，会标注）
    .xls   xlrd
    .md/.txt/.yaml  直接读
    .pdf   pypdf 抽文本层；抽不到说明是扫描件，提示去用 OCR
    图片   记成"素材"，不硬读（地图这种东西给 KP 一句"这里有张地图"比乱读有用）

**`.doc` 是启发式，不是解析器。** 老式 Word 二进制格式没有公开的稳定文本层接口，
这里是按 UTF-16LE 扫出连续可读片段再拼起来——文字基本都在，
但段落顺序、表格结构、页眉页脚可能不完美。所以会带一条 warning，
让守秘人知道这份材料是"读出来的"而不是"原本就是文本"。
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

# 能当正文读的
TEXT_SUFFIXES = {".md", ".txt", ".yaml", ".yml", ".csv", ".json"}
DOC_SUFFIXES = {".doc", ".docx", ".rtf"}
SHEET_SUFFIXES = {".xlsx", ".xlsm", ".xls"}
PDF_SUFFIXES = {".pdf"}
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp", ".tif", ".tiff"}

READABLE = (TEXT_SUFFIXES | DOC_SUFFIXES | SHEET_SUFFIXES
            | PDF_SUFFIXES | IMAGE_SUFFIXES)

_CJK = r"\u4e00-\u9fff\u3400-\u4dbf"
_PUNCT = r"\uff0c\u3002\uff01\uff1f\u3001\uff1a\uff1b\u201c\u201d\u2018\u2019\uff08\uff09\u300a\u300b\u3010\u3011\u2014\u2026"
_GOOD = re.compile(rf"[{_CJK}A-Za-z0-9{_PUNCT} \t]")


@dataclass
class DocResult:
    ok: bool = False
    text: str = ""
    kind: str = ""            # text | doc | sheet | pdf | image | unsupported
    warning: str = ""
    chars: int = 0
    meta: dict[str, Any] = field(default_factory=dict)


# ---------------------------------------------------------------- 各格式

def _read_text_file(p: Path) -> DocResult:
    for enc in ("utf-8", "utf-8-sig", "gbk", "big5", "latin-1"):
        try:
            t = p.read_text(encoding=enc)
        except Exception:
            continue
        if t.strip():
            if p.suffix.lower() in (".md", ".txt"):
                return DocResult(True, t.strip(), "text", chars=len(t))
            return DocResult(True, t.strip(), "text", chars=len(t))
    return DocResult(False, kind="text", warning="读不出内容（编码不认识或文件是空的）")


def _read_docx(p: Path) -> DocResult:
    """docx：能认出标题层级，输出带 `## ` 的文本，切场景就有依据。"""
    try:
        import docx
    except ImportError:
        return DocResult(False, kind="doc", warning="没装 python-docx，读不了 .docx")
    try:
        d = docx.Document(str(p))
    except Exception as e:  # noqa: BLE001
        # 常见情况：文件其实是宏启用的模板（.dotx/.docm 改后缀来的），
        # python-docx 会按内容类型拒收。那就自己解 zip 读 XML，
        # OOXML 这几种变体的正文结构是一样的。
        raw = _read_ooxml(p)
        if raw:
            return DocResult(True, raw, "doc", chars=len(raw),
                             warning="这个 .docx 其实是模板/宏格式，已改用直接读 XML 的方式提取。")
        return DocResult(False, kind="doc", warning=f"打不开：{type(e).__name__}: {e}")

    lines: list[str] = []
    for para in d.paragraphs:
        text = (para.text or "").strip()
        if not text:
            continue
        style = ""
        try:
            style = (para.style.name or "")
        except Exception:
            pass
        m = re.match(r"Heading\s*(\d+)", style, re.I)
        if m:
            lvl = max(2, min(3, int(m.group(1))))
            lines.append("#" * lvl + " " + text)
        elif style.lower().startswith("title"):
            lines.append("## " + text)
        else:
            lines.append(text)

    for ti, table in enumerate(d.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [(c.text or "").strip().replace("\n", " ") for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            lines.append(f"\n### 表 {ti}")
            lines.extend(rows)

    text = "\n".join(lines).strip()
    return DocResult(bool(text), text, "doc", chars=len(text),
                     warning="" if text else "这个 .docx 里没有文字")


def _read_ooxml(p: Path) -> str:
    """直接解 OOXML 包读正文，绕开 python-docx 的内容类型检查。

    这样 .docx / .dotx / .docm（以及改了后缀的）都能读——
    它们的正文都在 word/document.xml 里，结构完全一样。
    """
    import xml.etree.ElementTree as ET
    import zipfile

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        with zipfile.ZipFile(p) as z:
            target = next((n for n in z.namelist()
                           if n == "word/document.xml"
                           or n.endswith("/document.xml")), "")
            if not target:
                return ""
            xml = z.read(target)
    except Exception:
        return ""

    try:
        root = ET.fromstring(xml)
    except Exception:
        return ""

    def para_text(el: Any) -> str:
        return "".join(t.text or "" for t in el.iter(f"{W}t")).strip()

    def is_heading(el: Any) -> int:
        ppr = el.find(f"{W}pPr")
        if ppr is None:
            return 0
        st = ppr.find(f"{W}pStyle")
        if st is None:
            return 0
        val = (st.get(f"{W}val") or "").strip()
        m = re.match(r"Heading\s*(\d+)", val, re.I)
        if m:
            return max(2, min(3, int(m.group(1))))
        # 中文版 Word 的标题样式 id 常常就是 "1"/"2"/"3"
        if val in ("1", "2", "3"):
            return 2
        if re.match(r"^[aA]\d?$", val) and val.lower() in ("a1", "a2", "a3", "a"):
            return 2
        return 0

    lines: list[str] = []
    body = root.find(f"{W}body")
    if body is None:
        body = root
    for child in list(body):
        if child.tag == f"{W}p":
            text = para_text(child)
            if not text:
                continue
            lvl = is_heading(child)
            lines.append("#" * lvl + " " + text if lvl else text)
        elif child.tag == f"{W}tbl":
            rows = []
            for tr in child.iter(f"{W}tr"):
                cells = [para_text(tc) for tc in tr.iter(f"{W}tc")]
                if any(cells):
                    rows.append(" | ".join(cells))
            if rows:
                lines.append("\n### 表")
                lines.extend(rows)
    return "\n".join(lines).strip()


def _read_xlsx(p: Path) -> DocResult:
    try:
        import openpyxl
    except ImportError:
        return DocResult(False, kind="sheet", warning="没装 openpyxl，读不了 .xlsx")
    try:
        wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        return DocResult(False, kind="sheet", warning=f"打不开：{type(e).__name__}: {e}")
    parts: list[str] = []
    try:
        for ws in wb.worksheets:
            rows: list[str] = []
            for row in ws.iter_rows(values_only=True):
                cells = ["" if c is None else str(c).strip() for c in row]
                while cells and not cells[-1]:
                    cells.pop()
                if any(cells):
                    rows.append(" | ".join(cells))
                if len(rows) > 800:               # 表格别把上下文撑爆
                    rows.append("…（表格太长，后续略）")
                    break
            if rows:
                # 用三级标题：它只是表格内的小节，不该被当成切场景的依据
                parts.append(f"### 表：{ws.title}\n" + "\n".join(rows))
    finally:
        try:
            wb.close()
        except Exception:
            pass
    text = "\n\n".join(parts).strip()
    return DocResult(bool(text), text, "sheet", chars=len(text),
                     warning="" if text else "这个表格是空的")


def _read_xls(p: Path) -> DocResult:
    try:
        import xlrd
    except ImportError:
        return DocResult(False, kind="sheet", warning="没装 xlrd，读不了 .xls")
    try:
        book = xlrd.open_workbook(str(p))
    except Exception as e:  # noqa: BLE001
        return DocResult(False, kind="sheet", warning=f"打不开：{type(e).__name__}: {e}")
    parts: list[str] = []
    for ws in book.sheets():
        rows: list[str] = []
        for r in range(min(ws.nrows, 800)):
            cells = [str(ws.cell_value(r, c)).strip() for c in range(ws.ncols)]
            while cells and not cells[-1]:
                cells.pop()
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            parts.append(f"### 表：{ws.name}\n" + "\n".join(rows))
    text = "\n\n".join(parts).strip()
    return DocResult(bool(text), text, "sheet", chars=len(text),
                     warning="" if text else "这个表格是空的")


def _read_doc(p: Path) -> DocResult:
    """老式 .doc（OLE 二进制）：启发式扫文本。

    没有稳定的公开文本层接口，所以这里从 WordDocument 流里按 UTF-16LE
    扫出连续可读片段再拼。文字基本都在，但段落顺序和表格结构可能不完美——
    所以一定带 warning，让守秘人知道这是"读出来的"。
    """
    try:
        import olefile
    except ImportError:
        return DocResult(False, kind="doc", warning="没装 olefile，读不了 .doc")
    try:
        ole = olefile.OleFileIO(str(p))
    except Exception as e:  # noqa: BLE001
        return DocResult(False, kind="doc", warning=f"不是标准的 .doc 文件：{e}")

    best = ""
    try:
        if not ole.exists("WordDocument"):
            return DocResult(False, kind="doc", warning="里面没有 WordDocument 流")
        data = ole.openstream("WordDocument").read()
        for enc in ("utf-16-le", "gbk", "utf-8"):
            try:
                raw = data.decode(enc, errors="ignore")
            except Exception:
                continue
            cand = _extract_runs(raw)
            if len(cand) > len(best):
                best = cand
    finally:
        try:
            ole.close()
        except Exception:
            pass

    if not best.strip():
        return DocResult(False, kind="doc",
                         warning="读不出文字。可能是图片型文档，试试转成 PDF 走 OCR。")
    return DocResult(True, best, "doc", chars=len(best),
                     warning="这是从老式 .doc 里启发式扫出来的文字，"
                             "顺序和表格结构可能不完美，建议对照原文核对关键信息。")


def _extract_runs(raw: str) -> str:
    """从乱码流里挑出连续可读的文本段，并猜一下哪些像标题。"""
    runs = re.findall(
        rf"[{_CJK}A-Za-z0-9{_PUNCT} \t]{{4,}}", raw)
    lines: list[str] = []
    for r in runs:
        s = re.sub(r"\s{3,}", " ", r).strip()
        if len(s) < 4:
            continue
        # 太短的、没有标点结尾的，多半是小标题
        if 2 <= len(s) <= 24 and not re.search(r"[。！？，、：；]$", s):
            lines.append("## " + s)
        else:
            lines.append(s)
    # 去重相邻重复（二进制扫描容易重复吐同一段）
    out: list[str] = []
    for s in lines:
        if not out or out[-1] != s:
            out.append(s)
    return "\n".join(out)


def _read_pdf(p: Path) -> DocResult:
    try:
        import pypdf
    except ImportError:
        return DocResult(False, kind="pdf", warning="没装 pypdf，读不了 PDF")
    try:
        reader = pypdf.PdfReader(str(p))
        parts: list[str] = []
        for i, page in enumerate(reader.pages):
            try:
                parts.append(page.extract_text() or "")
            except Exception:
                continue
            if i > 400:
                parts.append("…（页数太多，后续略）")
                break
    except Exception as e:  # noqa: BLE001
        return DocResult(False, kind="pdf", warning=f"打不开：{type(e).__name__}: {e}")
    text = "\n\n".join(p.strip() for p in parts if p.strip()).strip()
    if not text:
        return DocResult(False, kind="pdf",
                         warning="抽不到文字层——这是扫描件，需要用 OCR 识别")
    return DocResult(True, text, "pdf", chars=len(text))


# ---------------------------------------------------------------- 统一入口

def read_any(path: Path) -> DocResult:
    p = Path(path)
    if not p.exists() or p.is_dir():
        return DocResult(False, warning="文件不存在")
    suffix = p.suffix.lower()
    if suffix in IMAGE_SUFFIXES:
        try:
            size = p.stat().st_size
        except OSError:
            size = 0
        return DocResult(False, kind="image", chars=0,
                         meta={"bytes": size},
                         warning="图片素材（地图/立绘之类），没有当文本读")
    if suffix == ".docx":
        return _read_docx(p)
    if suffix == ".doc":
        return _read_doc(p)
    if suffix in (".xlsx", ".xlsm"):
        return _read_xlsx(p)
    if suffix == ".xls":
        return _read_xls(p)
    if suffix == ".pdf":
        return _read_pdf(p)
    if suffix in TEXT_SUFFIXES:
        return _read_text_file(p)
    if suffix == ".rtf":
        r = _read_text_file(p)
        if r.ok:
            r.text = re.sub(r"\\[a-z]+-?\d* ?", "", r.text)
            r.warning = "这是 RTF，已粗略去掉了控制符"
        return r
    return DocResult(False, kind="unsupported", warning=f"不认识这种格式：{suffix}")


def scan_folder(root: Path) -> dict[str, Any]:
    """递归读一个模组文件夹，返回拼好的文本 + 素材清单 + 逐文件报告。"""
    parts: list[str] = []
    assets: list[str] = []
    skipped: list[str] = []
    reports: list[dict[str, Any]] = []
    warnings: list[str] = []

    files = sorted((f for f in Path(root).rglob("*") if f.is_file()),
                   key=lambda f: (len(f.relative_to(root).parts), str(f)))

    for f in files:
        rel = str(f.relative_to(root)).replace("\\", "/")
        if f.name.startswith(".") or f.name.startswith("~$"):
            continue                                  # Office 临时文件
        res = read_any(f)
        reports.append({"file": rel, "ok": res.ok, "kind": res.kind,
                        "chars": res.chars, "warning": res.warning})
        if res.ok and res.text.strip():
            parts.append(f"\n===== {rel} =====\n{res.text.strip()}")
        elif res.kind == "image":
            assets.append(rel)
        else:
            skipped.append(rel)
            if res.warning and res.kind not in ("image",):
                warnings.append(f"{rel}：{res.warning}")

    text = "\n".join(parts).strip()
    if assets:
        warnings.append("有 %d 个图片素材没有读（%s）——地图这类东西"
                        "给守秘人一句「这里有张地图」比乱读有用。"
                        % (len(assets), "、".join(assets[:4])))
    return {"text": text, "assets": assets, "skipped": skipped,
            "reports": reports, "warnings": warnings, "files": len(files)}
